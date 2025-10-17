from datetime import datetime
import json
from typing import List
import requests

import docx
import pptx
import fitz
import asyncio
from crawl4ai import AsyncWebCrawler, BrowserConfig, CacheMode, CrawlerRunConfig
from crawl4ai.extraction_strategy import JsonXPathExtractionStrategy

import logging


logger = logging.getLogger("iroko-cris")

def write_file(path: str, response: requests.Response) -> None:
    """
    Args:
    path (str): Path where the file will be saved
    response (requests.Response): requests.Response object
    """
    try:
        with open(path, "wb") as f:
            f.write(response.content)
    except FileNotFoundError as e:
        logger.error(e)
    except PermissionError  as e:
        logger.error(e)
    

def get_filename(url: str) -> str:
    filename = url.split("/")[-1].split(".")[0]
    return f"{filename}-{datetime.now().strftime('%Y%m%d%H%M%S%f')}"


async def start_extraction(start_urls: str) -> None:
    urls = [start_urls]
    visited = []

    while urls:
        current_url = urls.pop(0)
        visited.append(current_url)

        logger.info(f"Request to {current_url}")
        response = requests.get(current_url)
        try:
            content_type = response.headers.get("Content-Type", "")

            filename = get_filename(current_url)
            data = None
            match content_type:
                case "application/pdf" | "application/octet-stream":
                    path = f"uprchat/harvester/repositories/pdfs/{filename}.pdf"
                    write_file(
                        path,
                        response,
                    )
                    doc = fitz.open(path)
                    first_page = doc.load_page(0)
                    summary = first_page.get_text().encode("utf8")
                    data = {
                        "type": "document",
                        "url": current_url,
                        "summary": summary,
                        "stored_in": path
                    }
                    print(data)
                case (
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                    | "application/msword"
                    | "application/rtf"
                ):
                    path = f"uprchat/harvester/repositories/docs/{filename}.doc"
                    write_file(
                        path,
                        response,
                    )
                    f = open(path, "rb")
                    document = docx.Document(f)
                    i = 0
                    summary = ""
                    while i < len(document.paragraphs) and i < 30:
                        summary += f" {document.paragraphs[i].text}"
                        i += 1
                    data = {
                        "type": "document",
                        "url": current_url,
                        "summary": summary,
                        "stored_in": path
                    }
                    print(data)
                case (
                    "application/vnd.openxmlformats-officedocument.presentationml.presentation"
                    | "application/vnd.ms-powerpoint"
                    | "application/vnd.openxmlformats-officedocument.presentationml.slideshow"
                ):
                    path = f"uprchat/harvester/repositories/ppts/{filename}.ppt"
                    write_file(
                        path,
                        response,
                    )
                    f = open(path, "rb")
                    presentation = pptx.Presentation(f)
                    i = 0
                    summary = ""
                    while i < len(presentation.slides) and i < 30:
                        for shape in presentation.slides[i].shapes:
                            if shape.has_text_frame:
                                summary += f" {shape.text}"
                        i += 1
                    data = {
                        "type": "document",
                        "url": current_url,
                        "summary": summary,
                        "stored_in": path
                    }
                    print(data)
                case _:
                    result = await extraction_xpath_to_json(current_url)
                    if result is None or not result.success:
                        logger.error(f"Failed to retrieve information from {current_url}")
                        continue
                    logger.info(f"Successfully retrieved information from {current_url}")
                    content = json.loads(result.extracted_content)
                    data = {
                        "type": "page",
                        "url": current_url,
                        "title": content[0]["title"],
                        "body": content[0]["body"],
                    }
                    print(data)
                    internal_links = [
                        item["href"]
                        for item in result.links.get("internal", [])
                        if item["href"] not in visited and item["href"] not in urls
                    ]
                    urls.extend(internal_links)
        except requests.exceptions.RequestException as e:
            logger.error(e)
        except json.JSONDecodeError as e:
            logger.error(e)
        except fitz.FileDataError as e:
            logger.error(e)
        except docx.opc.exceptions.PackageNotFoundError as e:
            logger.error(e)
        except pptx.exc.PackageNotFoundError as e:
            logger.error(e)
        except IndexError | KeyError  as e:
            logger.error(e)


async def extraction_xpath_to_json(start_url: str):

    schema = {
        "name": "Data in plain text",
        "baseSelector": "/html",
        "fields": [
            {"name": "title", "selector": "//title", "type": "text"},
            {
                "name": "body",
                "selector": "//body//*[not(self::style or self::script)]",
                "type": "text",
            },
        ],
    }

    base_browser = BrowserConfig(
        headless=False,
        text_mode=True,
        user_agent_mode="random",
        java_script_enabled=True,
        accept_downloads=True,
        downloads_path="uprchat/harvester/repositories"
        # proxy_config={
        #     "server": "http://proxy.upr.edu.cu:8080",
        #     "username": "lazaro.hernandezp",
        #     "password": "LHP*2023",
        # }
    )

    config = CrawlerRunConfig(
        extraction_strategy=JsonXPathExtractionStrategy(schema, verbose=True),
        exclude_all_images=True,
        page_timeout=100000,
        cache_mode=CacheMode.BYPASS,
        wait_for="5000",
        js_code="""async () =>{
                    const downloadLinks = document.querySelectorAll('a[download]');
                    for (const link of downloadLinks) {
                        link.click();
                        // Delay between clicks
                        await new Promise(r => setTimeout(r, 2000));  
                    }
                }
            """
    )

    async with AsyncWebCrawler(config=base_browser) as crawler:
        result = await crawler.arun(url=start_url, config=config)
        return result


# asyncio.run(start_extraction(["https://arxiv.org/abs/2504.19667"]))
# asyncio.run(
#     start_extraction(
#         "https://cdn.www.gob.pe/uploads/document/file/3437716/1A.%20Modelo%20Informe%20Mu%CC%81ltiple.docx"
#     )
# )
# asyncio.run(
#     start_extraction(
#         "https://implanuruapan.gob.mx/wp-content/uploads/2021/01/Informe-Trimestral-Octubre-Diciembre_2020.pptx"
#     )
# )
asyncio.run(
    start_extraction(
        "https://www.upr.edu.cu/"
    )
)
